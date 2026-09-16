"""Text extraction from document bytes — shared by every knowledge source.

Both ingest paths (a Frappe File folder and a Google Drive folder) end up holding
the same thing: a filename and some bytes. Keeping one implementation means a PDF
indexes identically whichever folder it came from, and a parser fix lands in both
at once.

Every parser dependency is imported inside its own branch, on purpose. None of
them are hard requirements of the app: a site that only indexes markdown should
not need PyMuPDF installed, and a missing library degrades to "this one file was
skipped, here is why" rather than breaking the sync.
"""

import io

MAX_BYTES = 40 * 1024 * 1024      # a single knowledge document larger than this is a mistake

# What we will attempt to read. Anything else is skipped with a reason.
EXTRACTABLE_EXT = (".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md",
                   ".csv", ".json", ".yaml", ".yml")

MAX_SHEET_ROWS = 2000             # a spreadsheet beyond this is data, not knowledge


def extract_bytes(name, mime, blob):
    """(text, None) or (None, reason). Never raises."""
    if blob is None:
        return None, "no content"
    if len(blob) > MAX_BYTES:
        return None, "skipped: %.1f MB exceeds the %d MB limit" % (
            len(blob) / 1048576.0, MAX_BYTES // 1048576)

    low = (name or "").lower()
    mime = mime or ""

    if mime == "application/pdf" or low.endswith(".pdf"):
        return _pdf(blob)
    if low.endswith(".docx"):
        return _docx(blob)
    if low.endswith((".xlsx", ".xlsm")):
        return _xlsx(blob)
    if low.endswith((".txt", ".md", ".csv", ".json", ".yaml", ".yml")):
        return _text(blob)
    return None, "unsupported type %s" % (mime or low or "unknown")


def _pdf(blob):
    try:
        import fitz
    except ImportError:
        return None, "PDF support needs PyMuPDF (pip install pymupdf)"
    try:
        with fitz.open(stream=blob, filetype="pdf") as doc:
            parts = []
            for i in range(doc.page_count):
                t = doc.load_page(i).get_text("text").strip()
                if t:
                    parts.append(t)
        text = "\n\n".join(parts)
        # A scan has pages but no text layer. Saying so is useful; ingesting an
        # empty document silently is not.
        return (text, None) if text.strip() else (None, "PDF has no text layer (scan)")
    except Exception as e:
        return None, "PDF parse failed: %s" % str(e)[:120]


def _docx(blob):
    try:
        import docx
    except ImportError:
        return None, "docx support needs python-docx"
    try:
        d = docx.Document(io.BytesIO(blob))
        parts = [p.text for p in d.paragraphs if (p.text or "").strip()]
        # Tables carry the procedure steps in most SOPs, so they are not optional.
        for tbl in d.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts), None
    except Exception as e:
        return None, "docx parse failed: %s" % str(e)[:120]


def _xlsx(blob):
    try:
        import openpyxl
    except ImportError:
        return None, "xlsx support needs openpyxl"
    try:
        wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    break
                if any(c is not None for c in row):
                    rows.append(" | ".join("" if c is None else str(c) for c in row))
            if rows:
                # The sheet name becomes a heading, which the chunker splits on.
                parts.append("## %s\n%s" % (ws.title, "\n".join(rows)))
        wb.close()
        return "\n\n".join(parts), None
    except Exception as e:
        return None, "xlsx parse failed: %s" % str(e)[:120]


def _text(blob):
    try:
        import chardet
        enc = (chardet.detect(blob[:20000]) or {}).get("encoding") or "utf-8"
        return blob.decode(enc, errors="replace"), None
    except Exception:
        # chardet is a nicety; utf-8 with replacement is always better than failing.
        return blob.decode("utf-8", errors="replace"), None
